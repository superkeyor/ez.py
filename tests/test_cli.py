from contextlib import redirect_stderr, redirect_stdout
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

import pymupdf

from test_pptx2md import load_cli


class CliTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.namespace = load_cli()

    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="ez-cli-tests-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def pdf(self, name="paper.pdf", text="Highlighted words", annotated=True):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_text((72, 72), text)
            if annotated:
                page.add_highlight_annot(page.search_for(text))
            doc.set_toc([[1, "Section", 1]])
            doc.set_metadata({"producer": "Original producer", "title": "Original title"})
            doc.save(path)
        return path

    def run_cli(self, *arguments, command="pdfannot"):
        with redirect_stdout(io.StringIO()) as stdout, redirect_stderr(io.StringIO()) as stderr:
            try:
                status = self.namespace["_cli"]([str(arg) for arg in arguments], command=command)
            except SystemExit as exc:
                status = exc.code
        return status, stdout.getvalue(), stderr.getvalue()

    def marked(self, path):
        with pymupdf.open(path) as doc:
            self.assertEqual(doc.metadata["title"], "Original title")
            return "PyHighlightExtractor" in doc.metadata["producer"]

    def test_default_extracts_and_updates_metadata(self):
        source = self.pdf()
        status, output, errors = self.run_cli(source)
        self.assertEqual((status, errors), (0, ""))
        self.assertIn("Highlighted words (p. 1 Section)", output)
        self.assertIn("Highlights (out of 1 pages @", output)
        self.assertTrue(self.marked(source))

    def test_metadata_opt_out_leaves_source_bytes_unchanged(self):
        source = self.pdf("reader's paper.pdf")
        before = source.read_bytes()
        status, output, errors = self.run_cli(source, "--no-update-metadata")
        self.assertEqual((status, errors), (0, ""))
        self.assertIn("Highlighted words", output)
        self.assertEqual(source.read_bytes(), before)
        with redirect_stdout(io.StringIO()) as output:
            result = self.namespace["pdfannot"](source, update_metadata=False)
        self.assertIsNone(result)
        self.assertIn("Highlighted words", output.getvalue())
        self.assertEqual(source.read_bytes(), before)

    def test_explicit_metadata_update_and_already_marked_pdf(self):
        source = self.pdf()
        self.assertEqual(self.run_cli(source, "--update-metadata")[0], 0)
        marked = source.read_bytes()
        self.assertEqual(self.run_cli(source)[0], 0)
        self.assertEqual(source.read_bytes(), marked)

    def test_no_annotations_does_not_mark_pdf(self):
        source = self.pdf(annotated=False)
        before = source.read_bytes()
        self.assertEqual(self.run_cli(source), (0, "", ""))
        self.assertEqual(source.read_bytes(), before)

    def test_modes_and_timestamp_skip(self):
        source = self.pdf()
        for mode in (0, 1):
            with self.subTest(mode=mode):
                before = source.read_bytes()
                with patch("os.path.getctime", return_value=source.stat().st_mtime):
                    self.assertEqual(self.run_cli(source, "--mode", mode), (0, "", ""))
                self.assertEqual(source.read_bytes(), before)
        self.assertEqual(self.run_cli(source)[0], 0)
        with patch("os.path.getctime", return_value=source.stat().st_mtime - 120):
            self.assertEqual(self.run_cli(source, "--mode", 0), (0, "", ""))
            status, output, errors = self.run_cli(source, "--mode", 1)
        self.assertEqual((status, errors), (0, ""))
        self.assertIn("Highlighted words", output)

    def test_python_default_prepend_and_positional_mode(self):
        source = self.pdf()
        destination = self.root / "notes.md"
        destination.write_text("Existing notes", encoding="utf-8")
        with redirect_stdout(io.StringIO()) as output:
            result = self.namespace["pdfannot"](source, destination)
        self.assertIsNone(result)
        self.assertIn("Highlighted words", output.getvalue())
        self.assertTrue(destination.read_text().endswith("\nExisting notes"))
        self.assertTrue(self.marked(source))
        with redirect_stdout(io.StringIO()) as output:
            self.namespace["pdfannot"](source, None, 2, False)
        self.assertIn("Highlighted words", output.getvalue())
        with self.assertRaises(TypeError):
            self.namespace["pdfannot"](source, force=2)

    def test_invalid_modes_are_explicit_errors(self):
        source = self.pdf()
        self.assertEqual(self.run_cli(source, "--mode", 3)[0], 2)
        for mode in (-1, 3, True, "2"):
            with self.subTest(mode=mode), self.assertRaisesRegex(ValueError, "mode"):
                self.namespace["pdfannot"](source, mode=mode)
        self.assertFalse(self.marked(source))

    def test_metadata_updated_before_failed_output_write(self):
        source = self.pdf()
        with patch.dict(self.namespace, {"_output_write": Mock(
                side_effect=OSError("Output write failed"))}):
            status, output, errors = self.run_cli(source, "-o", self.root / "notes.md")
        self.assertEqual((status, output), (1, ""))
        self.assertIn("Output write failed", errors)
        self.assertTrue(self.marked(source))

    def test_combined_output_failure_reports_every_affected_input(self):
        first = self.pdf("a.pdf")
        second = self.pdf("b.pdf")
        destination = self.root / "notes.md"
        write = self.namespace["_output_write"]

        def fail_final_write(content, output, overwrite):
            if Path(output) == destination:
                self.assertTrue(self.marked(first))
                self.assertTrue(self.marked(second))
                raise OSError("Final output write failed")
            return write(content, output, overwrite)

        with patch.dict(self.namespace, {"_output_write": fail_final_write}):
            status, output, errors = self.run_cli(first, second, "-o", destination)
        self.assertEqual((status, output), (1, ""))
        self.assertIn("Final output write failed", errors)
        self.assertIn("0 succeeded, 2 failed", errors)
        self.assertFalse(destination.exists())

    def test_python_folder_preserves_prepend_order_and_continues_errors(self):
        first = self.pdf("a.pdf", "First annotation")
        second = self.pdf("b.pdf", "Second annotation")
        bad = self.root / "bad.pdf"
        bad.write_bytes(b"Not a PDF")
        timestamp = first.stat().st_mtime
        os.utime(first, (timestamp - 120, timestamp - 120))
        os.utime(second, (timestamp - 60, timestamp - 60))
        destination = self.root / "notes.md"
        destination.write_text("Old notes", encoding="utf-8")
        with patch.dict(self.namespace, {"fls": lambda path, regex: [second, bad, first]}):
            with redirect_stdout(io.StringIO()) as stdout, redirect_stderr(io.StringIO()) as stderr:
                self.namespace["pdfannot"](self.root, destination, update_metadata=False)
        content = destination.read_text()
        self.assertLess(content.index("Second annotation"), content.index("First annotation"))
        self.assertTrue(content.endswith("Old notes"))
        self.assertIn("First annotation", stdout.getvalue())
        self.assertIn("bad.pdf", stderr.getvalue())
        self.assertFalse(self.marked(first))
        self.assertFalse(self.marked(second))

    def test_metadata_failure_is_reported_and_other_inputs_continue(self):
        first = self.pdf("first.pdf")
        second = self.pdf("second.pdf", "Second highlight")
        destination = self.root / "notes.md"
        save = pymupdf.Document.save

        def fail_first(doc, filename, *args, **kwargs):
            if Path(filename) == first:
                raise OSError("Metadata write denied")
            return save(doc, filename, *args, **kwargs)

        with patch.object(pymupdf.Document, "save", fail_first):
            status, output, errors = self.run_cli(first, second, "-o", destination)
        self.assertEqual((status, output), (1, ""))
        self.assertIn("Cannot update PDF metadata", errors)
        self.assertIn("1 succeeded, 1 failed", errors)
        self.assertIn("Second highlight", destination.read_text())
        self.assertFalse(self.marked(first))
        self.assertTrue(self.marked(second))

    def test_combined_output_replaces_once_and_preserves_both_documents(self):
        first = self.pdf("a.pdf", "First annotation")
        second = self.pdf("b.pdf", "Second annotation")
        destination = self.root / "notes.md"
        destination.write_text("Old notes", encoding="utf-8")
        args = (first, second, "--output-file", destination, "--no-update-metadata")
        self.assertEqual(self.run_cli(*args)[0], 1)
        self.assertEqual(destination.read_text(), "Old notes")
        status, output, errors = self.run_cli(*args, "--overwrite")
        self.assertEqual((status, output), (0, ""))
        self.assertIn("2 succeeded, 0 failed", errors)
        content = destination.read_text()
        self.assertNotIn("Old notes", content)
        self.assertLess(content.index("# a.pdf"), content.index("# b.pdf"))
        self.assertIn("First annotation", content)
        self.assertIn("Second annotation", content)

    def test_failed_input_is_skipped_and_combined_output_is_written(self):
        first = self.pdf("a.pdf", "First annotation")
        bad = self.root / "bad.pdf"
        bad.write_bytes(b"Not a PDF")
        second = self.pdf("c.pdf", "Second annotation")
        destination = self.root / "notes.md"
        status, output, errors = self.run_cli(
            first, bad, second, "-o", destination, "--no-update-metadata")
        self.assertEqual((status, output), (1, ""))
        self.assertIn("bad.pdf", errors)
        self.assertIn("2 succeeded, 1 failed", errors)
        content = destination.read_text()
        self.assertIn("First annotation", content)
        self.assertIn("Second annotation", content)
        self.assertNotIn("# bad.pdf", content)

    def test_all_failed_inputs_leave_existing_destination_unchanged(self):
        bad = self.root / "bad.pdf"
        bad.write_bytes(b"Not a PDF")
        missing = self.root / "missing.pdf"
        destination = self.root / "notes.md"
        destination.write_text("Keep this", encoding="utf-8")
        status, _, errors = self.run_cli(bad, missing, "-o", destination, "--overwrite")
        self.assertEqual(status, 1)
        self.assertIn("0 succeeded, 2 failed", errors)
        self.assertEqual(destination.read_text(), "Keep this")
        destination.unlink()
        self.assertEqual(self.run_cli(bad, missing, "-o", destination)[0], 1)
        self.assertFalse(destination.exists())

    def test_empty_extractions_are_successful(self):
        first = self.pdf("empty-a.pdf", annotated=False)
        second = self.pdf("empty-b.pdf", annotated=False)
        destination = self.root / "notes.md"
        status, _, errors = self.run_cli(first, second, "-o", destination)
        self.assertEqual(status, 0)
        self.assertIn("2 succeeded, 0 failed", errors)
        self.assertIn("# empty-a.pdf", destination.read_text())
        self.assertFalse(self.marked(first))
        self.assertFalse(self.marked(second))

    def test_output_directory_recursion_and_duplicate_aliases(self):
        first = self.pdf("input/a.pdf", "First annotation")
        second = self.pdf("input/sub/b.PDF", "Second annotation")
        alias = self.root / "alias.pdf"
        alias.symlink_to(first)
        destination = self.root / "notes"
        status, output, errors = self.run_cli(
            self.root / "input", alias, first, "--recursive",
            "--output-dir", destination, "--no-update-metadata")
        self.assertEqual((status, output), (0, ""))
        self.assertIn("2 succeeded, 0 failed", errors)
        self.assertIn("First annotation", (destination / "input/a.md").read_text())
        self.assertIn("Second annotation", (destination / "input/sub/b.md").read_text())

    def test_wildcard_and_multiple_inputs_require_an_output_destination(self):
        self.pdf("a.pdf")
        self.pdf("b.pdf")
        status, _, errors = self.run_cli(self.root / "*.pdf")
        self.assertEqual(status, 2)
        self.assertIn("Multiple inputs require", errors)
        status, _, _ = self.run_cli(self.root / "*.pdf", "-o", self.root / "notes.md",
                                   "--no-update-metadata")
        self.assertEqual(status, 0)

    def test_duplicate_inputs_keep_single_document_format(self):
        source = self.pdf()
        destination = self.root / "notes.md"
        status, _, errors = self.run_cli(source, source, "-o", destination,
                                        "--no-update-metadata")
        self.assertEqual(status, 0, errors)
        self.assertIn("1 succeeded, 0 failed", errors)
        content = destination.read_text()
        self.assertTrue(content.startswith("### [paper.pdf]"))
        self.assertEqual(content.count("Highlighted words"), 1)

    def test_output_options_are_mutually_exclusive(self):
        source = self.pdf()
        status, _, errors = self.run_cli(
            source, "-o", self.root / "notes.md", "--output-dir", self.root / "out")
        self.assertEqual(status, 2)
        self.assertIn("mutually exclusive", errors)

    def test_source_and_aliases_cannot_be_output_destinations(self):
        source = self.pdf("a.pdf")
        second = self.pdf("b.pdf")
        symlink = self.root / "symlink.md"
        symlink.symlink_to(second)
        hardlink = self.root / "hardlink.md"
        os.link(second, hardlink)
        before = [source.read_bytes(), second.read_bytes()]
        for destination in (source, second, symlink, hardlink):
            with self.subTest(destination=destination):
                status, _, errors = self.run_cli(
                    source, second, "-o", destination, "--overwrite")
                self.assertEqual(status, 1)
                self.assertIn("Output must not replace", errors)
                self.assertEqual([source.read_bytes(), second.read_bytes()], before)

    def test_output_directory_collision_is_rejected(self):
        first = self.pdf("a/report.pdf")
        second = self.pdf("b/report.pdf")
        destination = self.root / "output"
        destination.mkdir()
        shared = self.root / "shared.md"
        shared.write_text("Keep", encoding="utf-8")
        for folder in ("a", "b"):
            (destination / folder).mkdir()
            os.link(shared, destination / folder / "report.md")
        status, _, errors = self.run_cli(
            first, second, "--output-dir", destination, "--overwrite")
        self.assertEqual(status, 1)
        self.assertIn("Output collision", errors)
        self.assertEqual(shared.read_text(), "Keep")

    def test_output_directory_escape_is_rejected(self):
        source = self.pdf("input/sub/paper.pdf")
        destination = self.root / "output"
        destination.mkdir()
        outside = self.root / "outside"
        outside.mkdir()
        (destination / "sub").symlink_to(outside, target_is_directory=True)
        status, _, errors = self.run_cli(
            source.parent.parent, "--recursive", "--output-dir", destination)
        self.assertEqual(status, 1)
        self.assertIn("escapes", errors)
        self.assertEqual(list(outside.iterdir()), [])

    def test_old_options_are_rejected_for_every_command(self):
        for command in self.namespace["_CLI_COMMANDS"]:
            for option in ("--output", "--force", "-f"):
                with self.subTest(command=command, option=option):
                    status, _, errors = self.run_cli("input.pdf", option, command=command)
                    self.assertEqual(status, 2)
                    self.assertIn("unrecognized arguments", errors)

    def make_document(self, command, name, text):
        if command in ("pdfannot", "pdf2md"):
            return self.pdf(name + ".pdf", text)
        if command == "pptx2md":
            from pptx import Presentation
            document = Presentation()
            document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = text
            path = self.root / (name + ".pptx")
        elif command == "docx2md":
            from docx import Document
            document = Document()
            document.add_paragraph(text)
            path = self.root / (name + ".docx")
        else:
            from openpyxl import Workbook
            document = Workbook()
            document.active.append(["Title"])
            document.active.append([text])
            path = self.root / (name + ".xlsx")
        document.save(path)
        return path

    def test_all_five_commands_share_combined_and_directory_output(self):
        for command in self.namespace["_CLI_COMMANDS"]:
            with self.subTest(command=command):
                first = self.make_document(command, command + "-first", "First document")
                second = self.make_document(command, command + "-second", "Second document")
                destination = self.root / (command + ".md")
                status, stdout, errors = self.run_cli(first, second, "-o", destination,
                                                     command=command)
                self.assertEqual((status, stdout), (0, ""), errors)
                content = destination.read_text()
                self.assertIn(f"# {first.name}", content)
                self.assertIn(f"# {second.name}", content)
                self.assertIn("First document", content)
                self.assertIn("Second document", content)
                directory = self.root / (command + "-out")
                status, stdout, errors = self.run_cli(
                    first, second, "--output-dir", directory, command=command)
                self.assertEqual((status, stdout), (0, ""), errors)
                self.assertIn("First document", (directory / (first.stem + ".md")).read_text())
                self.assertIn("Second document", (directory / (second.stem + ".md")).read_text())

    def test_shared_output_format_is_command_configured(self):
        first = self.root / "a.txt"
        second = self.root / "b.txt"
        first.write_text("First text", encoding="utf-8")
        second.write_text("Second text", encoding="utf-8")

        def copy_text(source, output=None, overwrite=False):
            return self.namespace["_output_write"](Path(source).read_text(), output, overwrite)

        spec = {
            "function": "copy_text", "output_format": "text",
            "description": "Copy text.", "arguments": [
                (["source"], {"nargs": "+", "help": "Input files."})],
            "groups": ["output", "batch"],
            "batch": {"input": "source", "extension": ".txt"},
            "notes": [], "examples": ["copy-text input.txt"],
        }
        with patch.dict(self.namespace, {"copy_text": copy_text,
                                        "combine_text": lambda results, root:
                                        "\n".join(content.rstrip("\n") for _, content in results)}):
            with patch.dict(self.namespace["_CLI_COMMANDS"], {"copy-text": spec}):
                with patch.dict(self.namespace["_CLI_OUTPUT_FORMATS"], {
                        "text": {"extension": ".txt", "combine": "combine_text"}}):
                    destination = self.root / "combined.txt"
                    status, _, errors = self.run_cli(first, second, "-o", destination,
                                                    command="copy-text")
                    self.assertEqual(status, 0, errors)
                    self.assertEqual(destination.read_text(), "First text\nSecond text\n")
                    directory = self.root / "out"
                    self.assertEqual(self.run_cli(first, second, "--output-dir", directory,
                                                  command="copy-text")[0], 0)
                    self.assertEqual((directory / "a.txt").read_text(), "First text\n")


if __name__ == "__main__":
    unittest.main()
