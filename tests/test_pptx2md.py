import ast
from contextlib import redirect_stderr, redirect_stdout
import io
import os
from pathlib import Path
import re
import sys
import tempfile
import unittest
import warnings
from types import SimpleNamespace
from unittest.mock import Mock, patch

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]


def load_cli():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", SyntaxWarning)
        tree = ast.parse((ROOT / "ez/ez.py").read_text(encoding="utf-8"))
    nodes = [
        node for node in tree.body
        if (isinstance(node, ast.FunctionDef)
            and (node.name.startswith(("_md_", "_output_", "_cli", "_pdfannot_"))
                 or node.name in ("_parse_ranges", "pptx2md", "docx2md", "xlsx2md",
                                  "pdf2md", "pdfannot")))
        or (isinstance(node, ast.Assign)
            and any(isinstance(target, ast.Name)
                    and target.id in ("_CLI_COMMANDS", "_CLI_ARGUMENT_GROUPS", "_CLI_OUTPUT_FORMATS")
                    for target in node.targets))
    ]
    import glob
    namespace = {"os": os, "re": re, "sys": sys, "glob": glob}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), "ez/ez.py", "exec"), namespace)
    return namespace


class Pptx2mdTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.namespace = load_cli()

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.source = Path(temporary.name) / "slides.pptx"
        self.presentation = Presentation()
        for title, visibility in (
            ("Visible first", None),
            ("Hidden zero", "0"),
            ("Hidden false", "false"),
            ("Visible one", "1"),
            ("Visible true", "true"),
        ):
            slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[0])
            slide.shapes.title.text = title
            slide.notes_slide.notes_text_frame.text = "Speaker note " + title
            if visibility is not None:
                slide._element.set("show", visibility)
        self.presentation.save(self.source)

    def convert(self, **kwargs):
        with redirect_stdout(io.StringIO()) as output:
            markdown = self.namespace["pptx2md"](self.source, **kwargs)
        self.assertEqual(output.getvalue(), markdown)
        return markdown

    def add_image_to_first_slide(self):
        from PIL import Image

        image_path = self.source.with_suffix(".png")
        Image.new("RGB", (32, 32), "white").save(image_path)
        self.presentation.slides[0].shapes.add_picture(
            str(image_path), Inches(1), Inches(1)
        )
        self.presentation.save(self.source)

    def test_default_skips_hidden_and_preserves_original_numbers(self):
        markdown = self.convert()
        self.assertEqual(re.findall(r"Slide number: (\d+)", markdown), ["1", "4", "5"])
        self.assertNotIn("Hidden", markdown)
        self.assertNotIn("Speaker note", markdown)
        self.assertIn("Visible first", markdown)
        self.assertIn("Visible one", markdown)
        self.assertIn("Visible true", markdown)

    def test_include_hidden_with_notes(self):
        markdown = self.convert(skip_hidden=False, note=True)
        self.assertEqual(re.findall(r"Slide number: (\d+)", markdown),
                         ["1", "2", "3", "4", "5"])
        self.assertIn("Hidden zero", markdown)
        self.assertIn("Hidden false", markdown)
        self.assertIn("Speaker note Hidden zero", markdown)

    def test_selection_uses_original_numbers_and_skips_hidden(self):
        markdown = self.convert(slides="4,2,1,4")
        self.assertEqual(re.findall(r"Slide number: (\d+)", markdown), ["1", "4"])
        self.assertNotIn("Hidden", markdown)
        markdown = self.convert(slides="2-3", skip_hidden=False)
        self.assertEqual(re.findall(r"Slide number: (\d+)", markdown), ["2", "3"])

    def test_hidden_only_selection_is_empty(self):
        self.assertEqual(self.convert(slides="2-3"), "")

    def test_all_hidden_presentation_is_empty(self):
        for slide in self.presentation.slides:
            slide._element.set("show", "0")
        self.presentation.save(self.source)
        self.assertEqual(self.convert(), "")

    def test_without_markers(self):
        markdown = self.convert(markers=False)
        self.assertNotIn("Slide number:", markdown)
        self.assertNotIn("Hidden", markdown)
        self.assertEqual(markdown.count("\n\n---\n\n"), 2)

    def test_ocr_is_opt_in_and_text_is_added_to_its_slide(self):
        self.add_image_to_first_slide()
        recognize = Mock(return_value="Recognized picture text")
        failed_dir = self.source.parent / "failed-images"
        with patch.dict(self.namespace, {
                "_md_require_vision": Mock(),
                "_md_vision_ocr": recognize,
        }):
            default_markdown = self.convert()
            self.assertNotIn("Recognized picture text", default_markdown)
            recognize.assert_not_called()

            markdown = self.convert(
                ocr=True, ocr_failed_dir=failed_dir, overwrite=True
            )

        self.assertIn("<!-- Slide number: 1 -->", markdown)
        self.assertIn("*[Image OCR]*\n\nRecognized picture text\n\n*[End OCR]*", markdown)
        recognize.assert_called_once()
        self.assertEqual(recognize.call_args.args[1], "accurate")
        self.assertEqual(recognize.call_args.args[2], "english")
        self.assertEqual(recognize.call_args.kwargs["failed_dir"], str(failed_dir))
        self.assertTrue(recognize.call_args.kwargs["overwrite"])
        self.assertIn(str(self.source), recognize.call_args.args[3])
        self.assertFalse(failed_dir.exists())

    def test_pptx_ocr_respects_picture_crop_and_later_opaque_cover(self):
        from PIL import Image
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

        image_path = self.source.with_suffix(".png")
        Image.new("RGB", (100, 80), "white").save(image_path)
        slide = self.presentation.slides[0]
        picture = slide.shapes.add_picture(
            str(image_path), Inches(1), Inches(1), Inches(4), Inches(2)
        )
        picture.crop_left = 0.1
        picture.crop_right = 0.2
        cover = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            picture.left + picture.width // 2,
            picture.top,
            picture.width // 4,
            picture.height,
        )
        cover.fill.solid()
        cover.fill.fore_color.rgb = RGBColor(255, 255, 255)

        image_data, name, extension = next(
            iter(self.namespace["_md_pptx_image_blobs"](slide))
        )
        with Image.open(io.BytesIO(image_data)) as visible_image:
            self.assertEqual(visible_image.size, (70, 80))
            self.assertEqual(visible_image.getpixel((5, 40)), (255, 255, 255, 255))
            # The frame shows the cropped image; the cover spans 2-3 in of 4 in.
            self.assertEqual(visible_image.getpixel((30, 40)), (255, 255, 255, 255))
            self.assertEqual(visible_image.getpixel((36, 40)), (0, 0, 0, 255))
            self.assertEqual(visible_image.getpixel((52, 40)), (0, 0, 0, 255))
            self.assertEqual(visible_image.getpixel((54, 40)), (255, 255, 255, 255))
        self.assertEqual(name, picture.name)
        self.assertEqual(extension, "png")

    def test_pptx_ocr_skips_fully_covered_picture(self):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

        self.add_image_to_first_slide()
        slide = self.presentation.slides[0]
        picture = next(
            shape for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        cover = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            picture.left,
            picture.top,
            picture.width,
            picture.height,
        )
        cover.fill.solid()
        cover.fill.fore_color.rgb = RGBColor(255, 255, 255)
        self.presentation.save(self.source)
        recognize = Mock()
        with patch.dict(self.namespace, {
                "_md_check_ocr_platform": Mock(),
                "_md_vision_ocr": recognize,
        }):
            self.convert(ocr=True)
        recognize.assert_not_called()

    def test_pptx_ocr_does_not_mask_transparent_cover_shape(self):
        from PIL import Image
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.oxml.xmlchemy import OxmlElement

        image_path = self.source.with_suffix(".png")
        Image.new("RGB", (32, 32), "white").save(image_path)
        slide = self.presentation.slides[0]
        picture = slide.shapes.add_picture(
            str(image_path), Inches(1), Inches(1), Inches(2), Inches(2)
        )
        cover = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            picture.left,
            picture.top,
            picture.width,
            picture.height,
        )
        cover.fill.solid()
        cover.fill.fore_color.rgb = RGBColor(255, 255, 255)
        alpha = OxmlElement("a:alpha")
        alpha.set("val", "50000")
        cover._element.spPr.solidFill.srgbClr.append(alpha)

        image_data, _, _ = next(
            iter(self.namespace["_md_pptx_image_blobs"](slide))
        )
        self.assertEqual(image_data, picture.image.blob)

    def test_pptx_ocr_masks_only_axis_aligned_rectangles(self):
        from PIL import Image
        from pptx.enum.shapes import MSO_SHAPE

        for cover_type, cover_rotation, picture_rotation in (
                (MSO_SHAPE.OVAL, 0, 0),
                (MSO_SHAPE.RECTANGLE, 15, 0),
                (MSO_SHAPE.RECTANGLE, 0, 15)):
            with self.subTest(
                    cover_type=cover_type,
                    cover_rotation=cover_rotation,
                    picture_rotation=picture_rotation):
                image_path = self.source.with_suffix(".png")
                Image.new("RGB", (32, 32), "white").save(image_path)
                slide = self.presentation.slides.add_slide(
                    self.presentation.slide_layouts[6]
                )
                picture = slide.shapes.add_picture(
                    str(image_path), Inches(1), Inches(1), Inches(2), Inches(2)
                )
                picture.rotation = picture_rotation
                cover = slide.shapes.add_shape(
                    cover_type,
                    picture.left,
                    picture.top,
                    picture.width,
                    picture.height,
                )
                cover.fill.solid()
                cover.rotation = cover_rotation

                image_data, _, _ = next(
                    iter(self.namespace["_md_pptx_image_blobs"](slide))
                )
                self.assertEqual(image_data, picture.image.blob)

    def add_picture_with_cover(self, slide, cover_geometry="roundRect", radius="0"):
        """Add a 32x32 white picture and a white cover over its left half."""
        from PIL import Image
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        image_path = self.source.with_suffix(".png")
        Image.new("RGB", (32, 32), "white").save(image_path)
        picture = slide.shapes.add_picture(
            str(image_path), Inches(1), Inches(1), Inches(2), Inches(2)
        )
        cover = slide.shapes.add_textbox(
            picture.left, picture.top, picture.width // 2, picture.height
        )
        cover.fill.solid()
        cover.fill.fore_color.rgb = RGBColor(255, 255, 255)
        geometry = cover._element.spPr.get_or_add_prstGeom()
        geometry.set("prst", cover_geometry)
        if radius is not None:
            from pptx.oxml import parse_xml
            geometry.append(parse_xml(
                '<a:avLst xmlns:a="http://schemas.openxmlformats.org/'
                'drawingml/2006/main"><a:gd name="adj" fmla="val %s"/></a:avLst>'
                % radius
            ))
        return picture, cover

    def first_blob_pixels(self, slide):
        from PIL import Image

        image_data, _, _ = next(iter(self.namespace["_md_pptx_image_blobs"](slide)))
        with Image.open(io.BytesIO(image_data)) as image:
            return image.getpixel((4, 16)), image.getpixel((28, 16))

    def test_pptx_ocr_masks_square_round_rect_text_box_cover(self):
        slide = self.presentation.slides[0]
        self.add_picture_with_cover(slide, "roundRect", "0")
        self.assertEqual(
            self.first_blob_pixels(slide), ((0, 0, 0, 255), (255, 255, 255, 255))
        )

    def test_pptx_ocr_ignores_rounded_round_rect_cover(self):
        slide = self.presentation.slides[0]
        picture, _ = self.add_picture_with_cover(slide, "roundRect", "16667")
        image_data, _, _ = next(iter(self.namespace["_md_pptx_image_blobs"](slide)))
        self.assertEqual(image_data, picture.image.blob)

    def test_pptx_ocr_ignores_hidden_cover_and_skips_hidden_picture(self):
        slide = self.presentation.slides[0]
        picture, cover = self.add_picture_with_cover(slide)
        cover._element.xpath("./p:nvSpPr/p:cNvPr")[0].set("hidden", "1")
        image_data, _, _ = next(iter(self.namespace["_md_pptx_image_blobs"](slide)))
        self.assertEqual(image_data, picture.image.blob)
        picture._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("hidden", "1")
        self.assertEqual(list(self.namespace["_md_pptx_image_blobs"](slide)), [])

    def test_pptx_ocr_resolves_cover_in_scaled_group_and_later_top_level_cover(self):
        from pptx.dml.color import RGBColor

        slide = self.presentation.slides[0]
        picture, cover = self.add_picture_with_cover(slide)
        # Move picture and cover into a group that scales child space 2x.
        group = slide.shapes.add_group_shape([picture, cover])
        xfrm = group._element.grpSpPr.xfrm
        xfrm.ext.cx *= 2
        xfrm.ext.cy *= 2
        self.assertEqual(
            self.first_blob_pixels(slide), ((0, 0, 0, 255), (255, 255, 255, 255))
        )
        # A later top-level cover over the right quarter of the grouped picture.
        late = slide.shapes.add_shape(
            1, group.left + group.width * 3 // 4, group.top,
            group.width // 4, group.height,
        )
        late.fill.solid()
        late.fill.fore_color.rgb = RGBColor(255, 255, 255)
        self.assertEqual(
            self.first_blob_pixels(slide), ((0, 0, 0, 255), (0, 0, 0, 255))
        )

    def test_pptx_ocr_skips_movies(self):
        from PIL import Image

        image_path = self.source.with_suffix(".png")
        Image.new("RGB", (32, 32), "white").save(image_path)
        video_path = self.source.with_suffix(".mp4")
        video_path.write_bytes(b"\x00\x00\x00\x18ftypmp42")
        slide = self.presentation.slides[0]
        slide.shapes.add_movie(
            str(video_path), Inches(1), Inches(1), Inches(2), Inches(2),
            poster_frame_image=str(image_path),
        )
        self.assertEqual(list(self.namespace["_md_pptx_image_blobs"](slide)), [])

    def test_pptx_ocr_includes_picture_placeholders(self):
        from PIL import Image

        image_path = self.source.with_suffix(".png")
        Image.new("RGB", (32, 32), "white").save(image_path)
        slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[8])
        placeholder = next(
            shape for shape in slide.placeholders
            if shape.placeholder_format.type == 18
        )
        placeholder.insert_picture(str(image_path))
        blobs = list(self.namespace["_md_pptx_image_blobs"](slide))
        self.assertEqual(len(blobs), 1)

    def test_pptx_ocr_skips_unrenderable_images_when_filtering_is_needed(self):
        from PIL import Image
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        image_path = self.source.with_suffix(".png")
        Image.new("RGB", (32, 32), "white").save(image_path)
        slide = self.presentation.slides[0]
        picture = slide.shapes.add_picture(
            str(image_path), Inches(1), Inches(1), Inches(2), Inches(2)
        )
        cover = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            picture.left,
            picture.top,
            picture.width // 2,
            picture.height,
        )
        cover.fill.solid()
        cover.fill.fore_color.rgb = RGBColor(255, 255, 255)

        with (patch.dict(self.namespace, {
                "_md_rasterize_emf_pdf": Mock(return_value=None),
        }), patch("PIL.Image.open", side_effect=OSError("unsupported image")),
              warnings.catch_warnings(record=True) as caught):
            warnings.simplefilter("always")
            blobs = list(self.namespace["_md_pptx_image_blobs"](slide))

        self.assertEqual(blobs, [])
        self.assertEqual(len(caught), 1)
        self.assertIn("Skipping OCR for PowerPoint picture", str(caught[0].message))
        self.assertIn("cannot apply crop or cover filtering", str(caught[0].message))

    def test_ocr_cli_uses_requested_mode(self):
        self.add_image_to_first_slide()
        recognize = Mock(return_value="Fast recognized text")
        failed_dir = self.source.parent / "failed-images"
        with (patch.dict(self.namespace, {
                "_md_require_vision": Mock(),
                "_md_vision_ocr": recognize,
        }), redirect_stdout(io.StringIO()) as output):
            status = self.namespace["_cli"](
                [str(self.source), "--ocr", "--ocr-mode", "fast",
                 "--ocr-failed-dir", str(failed_dir), "--overwrite"],
                command="pptx2md",
            )
        self.assertEqual(status, 0)
        self.assertIn("Fast recognized text", output.getvalue())
        self.assertEqual(recognize.call_args.args[1], "fast")
        self.assertEqual(recognize.call_args.kwargs["failed_dir"], str(failed_dir))
        self.assertTrue(recognize.call_args.kwargs["overwrite"])

    def test_ocr_cli_uses_requested_language(self):
        self.add_image_to_first_slide()
        recognize = Mock(return_value="Chinese recognized text")
        with (patch.dict(self.namespace, {
                "_md_require_vision": Mock(),
                "_md_vision_ocr": recognize,
        }), redirect_stdout(io.StringIO())):
            status = self.namespace["_cli"](
                [str(self.source), "--ocr", "--ocr-language", "chinese"],
                command="pptx2md",
            )
        self.assertEqual(status, 0)
        self.assertEqual(recognize.call_args.args[2], "chinese")

    def test_ocr_mode_requires_ocr_on_cli(self):
        with redirect_stderr(io.StringIO()) as errors:
            with self.assertRaises(SystemExit) as result:
                self.namespace["_cli"](
                    [str(self.source), "--ocr-mode", "fast"], command="pptx2md"
                )
        self.assertEqual(result.exception.code, 2)
        self.assertIn("require --ocr", errors.getvalue())

    def test_ocr_language_requires_ocr_on_cli(self):
        with redirect_stderr(io.StringIO()) as errors:
            with self.assertRaises(SystemExit) as result:
                self.namespace["_cli"](
                    [str(self.source), "--ocr-language", "chinese"],
                    command="pptx2md",
                )
        self.assertEqual(result.exception.code, 2)
        self.assertIn("--ocr-language", errors.getvalue())

    def test_ocr_failed_dir_requires_ocr_on_cli(self):
        with redirect_stderr(io.StringIO()) as errors:
            with self.assertRaises(SystemExit) as result:
                self.namespace["_cli"](
                    [str(self.source), "--ocr-failed-dir", "failed-images"],
                    command="pptx2md",
                )
        self.assertEqual(result.exception.code, 2)
        self.assertIn("--ocr-failed-dir", errors.getvalue())

    def test_ocr_failed_dir_requires_ocr_in_python_api(self):
        with self.assertRaisesRegex(ValueError, "requires ocr=True"):
            self.convert(ocr_failed_dir=self.source.parent / "failed-images")

    def test_vision_ocr_rejects_non_macos(self):
        with patch.object(sys, "platform", "linux"):
            with self.assertRaisesRegex(OSError, "only on macOS"):
                self.namespace["_md_require_vision"]()

    def test_vision_ocr_uses_requested_level_and_languages(self):
        from PIL import Image

        image = io.BytesIO()
        Image.new("RGB", (4, 4), "white").save(image, format="PNG")
        image_data = image.getvalue()
        for mode, level in (("accurate", 1), ("fast", 2)):
            for language, languages in (
                    ("english", ["en-US"]),
                    ("chinese", ["zh-Hans"]),
                    ("both", ["zh-Hans", "en-US"])):
                request = Mock()
                request.results.return_value = []
                handler = Mock()
                handler.performRequests_error_.return_value = (True, None)
                vision = SimpleNamespace(
                    VNRecognizeTextRequest=Mock(),
                    VNRequestTextRecognitionLevelAccurate=1,
                    VNRequestTextRecognitionLevelFast=2,
                    VNImageRequestHandler=Mock(),
                )
                vision.VNRecognizeTextRequest.alloc.return_value.init.return_value = request
                vision.VNImageRequestHandler.alloc.return_value.initWithData_options_.return_value = handler
                foundation = SimpleNamespace(
                    NSData=SimpleNamespace(
                        dataWithBytes_length_=Mock(return_value=object())
                    )
                )
                with (patch.object(sys, "platform", "darwin"),
                      patch.dict(sys.modules, {
                          "Vision": vision, "Foundation": foundation
                      })):
                    self.assertEqual(
                        self.namespace["_md_vision_ocr"](image_data, mode, language),
                        "",
                    )
                request.setRecognitionLevel_.assert_called_once_with(level)
                request.setRecognitionLanguages_.assert_called_once_with(languages)

    def test_vision_ocr_reports_missing_binding(self):
        with (patch.object(sys, "platform", "darwin"),
              patch.dict(sys.modules, {"Vision": None})):
            with self.assertRaisesRegex(ImportError, "pyobjc-framework-Vision"):
                self.namespace["_md_require_vision"]()

    def test_vision_ocr_skips_invalid_image_with_warning(self):
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            result = self.namespace["_md_vision_ocr"](
                b"not an image", context="PowerPoint slide 2"
            )
        self.assertEqual(result, "")
        self.assertEqual(len(caught), 1)
        self.assertIn("PowerPoint slide 2", str(caught[0].message))

    def test_vision_zero_dimension_error_skips_only_that_image(self):
        request = Mock()
        request.results.return_value = []
        handler = Mock()
        handler.performRequests_error_.return_value = (
            False, "Error: zero-dimensioned image (0 x 0)"
        )
        vision = SimpleNamespace(
            VNRecognizeTextRequest=Mock(),
            VNRequestTextRecognitionLevelAccurate=1,
            VNRequestTextRecognitionLevelFast=2,
            VNImageRequestHandler=Mock(),
        )
        vision.VNRecognizeTextRequest.alloc.return_value.init.return_value = request
        vision.VNImageRequestHandler.alloc.return_value.initWithData_options_.return_value = handler
        foundation = SimpleNamespace(
            NSData=SimpleNamespace(dataWithBytes_length_=Mock(return_value=object()))
        )
        image = io.BytesIO()
        from PIL import Image

        Image.new("RGB", (4, 4), "white").save(image, format="PNG")
        failed_dir = self.source.parent / "failed-images"
        with (patch.object(sys, "platform", "darwin"),
              patch.dict(sys.modules, {"Vision": vision, "Foundation": foundation}),
              warnings.catch_warnings(record=True) as caught):
            warnings.simplefilter("always")
            result = self.namespace["_md_vision_ocr"](
                image.getvalue(), context="PowerPoint slide 3",
                failed_dir=failed_dir,
                failed_name="deck-slide-3-Graphic 1",
                image_extension=".png",
            )
        self.assertEqual(result, "")
        self.assertEqual(len(caught), 1)
        self.assertIn("zero-dimensioned", str(caught[0].message))
        saved_image = failed_dir / "deck-slide-3-Graphic_1.png"
        self.assertEqual(saved_image.read_bytes(), image.getvalue())
        self.assertIn(str(saved_image), str(caught[0].message))

    def test_vision_failure_saves_image_before_raising(self):
        from PIL import Image

        request = Mock()
        handler = Mock()
        handler.performRequests_error_.return_value = (
            False, "Vision processing failed"
        )
        vision = SimpleNamespace(
            VNRecognizeTextRequest=Mock(),
            VNRequestTextRecognitionLevelAccurate=1,
            VNRequestTextRecognitionLevelFast=2,
            VNImageRequestHandler=Mock(),
        )
        vision.VNRecognizeTextRequest.alloc.return_value.init.return_value = request
        vision.VNImageRequestHandler.alloc.return_value.initWithData_options_.return_value = handler
        foundation = SimpleNamespace(
            NSData=SimpleNamespace(dataWithBytes_length_=Mock(return_value=object()))
        )
        image = io.BytesIO()
        Image.new("RGB", (4, 4), "white").save(image, format="PNG")
        image_data = image.getvalue()
        failed_dir = self.source.parent / "failed-images"
        with (patch.object(sys, "platform", "darwin"),
              patch.dict(sys.modules, {"Vision": vision, "Foundation": foundation})):
            with self.assertRaisesRegex(OSError, "Apple Vision OCR failed"):
                self.namespace["_md_vision_ocr"](
                    image_data,
                    context="PowerPoint slide 4, picture 'Graphic 1'",
                    failed_dir=failed_dir,
                    failed_name="deck-slide-4-Graphic 1",
                    image_extension=".png",
                )

        saved_image = failed_dir / "deck-slide-4-Graphic_1.png"
        self.assertEqual(saved_image.read_bytes(), image_data)

    def test_failed_ocr_image_saves_without_overwriting_existing_files(self):
        failed_dir = self.source.parent / "failed-images"
        save_failed = self.namespace["_md_save_failed_ocr_image"]

        first = save_failed(b"first", failed_dir, "slide/1", ".PNG")
        second = save_failed(b"second", failed_dir, "slide/1", ".PNG")

        self.assertEqual(first.name, "slide_1.png")
        self.assertEqual(second.name, "slide_1-2.png")
        self.assertEqual(first.read_bytes(), b"first")
        self.assertEqual(second.read_bytes(), b"second")

    def test_failed_ocr_image_overwrite_replaces_same_named_file(self):
        failed_dir = self.source.parent / "failed-images"
        save_failed = self.namespace["_md_save_failed_ocr_image"]

        first = save_failed(b"first", failed_dir, "slide/1", ".PNG")
        replaced = save_failed(
            b"replacement", failed_dir, "slide/1", ".PNG", overwrite=True
        )

        self.assertEqual(first, replaced)
        self.assertEqual(first.name, "slide_1.png")
        self.assertEqual(first.read_bytes(), b"replacement")

    def test_pptx_with_no_embedded_images_still_converts(self):
        with (patch.dict(self.namespace, {"_md_check_ocr_platform": Mock()}),
              redirect_stdout(io.StringIO()) as output):
            markdown = self.namespace["pptx2md"](self.source, ocr=True)
        self.assertEqual(output.getvalue(), markdown)
        self.assertIn("Visible first", markdown)
        self.assertNotIn("Image OCR", markdown)

    def test_pptx_linked_picture_is_silently_skipped(self):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        from pptx.oxml.ns import qn

        self.add_image_to_first_slide()
        slide = self.presentation.slides[0]
        picture = next(
            shape for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        picture._pic.blipFill.blip.attrib.pop(qn("r:embed"))
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            blobs = list(self.namespace["_md_pptx_image_blobs"](slide))
        self.assertEqual(blobs, [])
        self.assertEqual(caught, [])

    def test_pptx_linked_picture_does_not_fail_conversion(self):
        from pptx.oxml.ns import qn
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        self.add_image_to_first_slide()
        picture = next(
            shape for shape in self.presentation.slides[0].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        picture._pic.blipFill.blip.attrib.pop(qn("r:embed"))
        self.presentation.save(self.source)
        recognize = Mock()
        with (patch.dict(self.namespace, {
                "_md_check_ocr_platform": Mock(),
                "_md_vision_ocr": recognize,
        }), warnings.catch_warnings(record=True) as caught):
            warnings.simplefilter("always")
            markdown = self.convert(ocr=True)
        self.assertIn("Visible first", markdown)
        self.assertEqual(caught, [])
        recognize.assert_not_called()

    def test_cli_pptx_linked_picture_produces_no_warning(self):
        from pptx.oxml.ns import qn
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        self.add_image_to_first_slide()
        picture = next(
            shape for shape in self.presentation.slides[0].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        picture._pic.blipFill.blip.attrib.pop(qn("r:embed"))
        self.presentation.save(self.source)
        errors = io.StringIO()
        recognize = Mock()
        with (patch.dict(self.namespace, {
                "_md_check_ocr_platform": Mock(),
                "_md_vision_ocr": recognize,
        }),
              redirect_stderr(errors), redirect_stdout(io.StringIO())):
            status = self.namespace["_cli"](
                [str(self.source), "--ocr"], command="pptx2md"
            )

        self.assertEqual(status, 0)
        self.assertEqual(errors.getvalue(), "")
        recognize.assert_not_called()

    def test_cli_pptx_vision_warning_names_file_slide_and_picture(self):
        self.add_image_to_first_slide()
        self.presentation.save(self.source)

        def skip_vision(image_data, mode, language, context, **kwargs):
            return skip_vision_with_warning(context)

        def skip_vision_with_warning(context):
            warnings.warn(
                f"Skipping OCR for {context}: Vision reports a zero-dimensioned image.",
                UserWarning,
            )
            return ""

        errors = io.StringIO()
        with (patch.dict(self.namespace, {
                "_md_check_ocr_platform": Mock(),
                "_md_vision_ocr": skip_vision,
        }), redirect_stderr(errors), redirect_stdout(io.StringIO())):
            status = self.namespace["_cli"](
                [str(self.source), "--ocr"], command="pptx2md"
            )

        self.assertEqual(status, 0)
        warning = errors.getvalue()
        self.assertIn("pptx2md: warning:", warning)
        self.assertIn(str(self.source), warning)
        self.assertIn("slide 1", warning)
        self.assertIn("picture 'Picture 3'", warning)
        self.assertIn("zero-dimensioned image", warning)
        self.assertNotRegex(warning, r"ez\.py:\d+")

    def test_cli_default_and_include_hidden(self):
        for flags, expected in (([], ["1", "4", "5"]),
                                (["--include-hidden"], ["1", "2", "3", "4", "5"])):
            with self.subTest(flags=flags), redirect_stdout(io.StringIO()) as output:
                status = self.namespace["_cli"]([str(self.source), *flags], command="pptx2md")
            self.assertEqual(status, 0)
            self.assertEqual(re.findall(r"Slide number: (\d+)", output.getvalue()), expected)

    def test_cli_help_for_all_converters(self):
        for command in self.namespace["_CLI_COMMANDS"]:
            for flag in ("-h", "--help"):
                with self.subTest(command=command, flag=flag):
                    with redirect_stdout(io.StringIO()) as output:
                        with self.assertRaises(SystemExit) as result:
                            self.namespace["_cli"]([flag], command=command)
                    self.assertEqual(result.exception.code, 0)
                    help_text = output.getvalue()
                    self.assertIn(f"Usage:\n    {command} ", help_text)
                    for section in ("Notes:", "Example:", "See also:"):
                        self.assertIn(section, help_text)
                    if command == "pptx2md":
                        self.assertIn("--include-hidden", help_text)
                        self.assertIn("skipped by default", help_text)
                        self.assertIn("Slides are 1-based", help_text)
                        self.assertIn("original order and numbering", help_text)
                        self.assertIn("even when selected", help_text)
                    if command in ("pptx2md", "docx2md", "xlsx2md"):
                        self.assertIn("--ocr", help_text)
                        self.assertIn("--ocr-mode", help_text)
                        self.assertIn("--ocr-language", help_text)
                        self.assertIn("--ocr-failed-dir", help_text)

    def test_cli_without_arguments_shows_help(self):
        for command in self.namespace["_CLI_COMMANDS"]:
            with self.subTest(command=command), redirect_stdout(io.StringIO()) as output:
                status = self.namespace["_cli"]([], command=command)
            self.assertEqual(status, 0)
            self.assertIn(f"Usage:\n    {command} ", output.getvalue())


class OfficeOcrTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.namespace = load_cli()

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        from PIL import Image

        self.image_path = self.root / "image.png"
        Image.new("RGB", (32, 32), "white").save(self.image_path)

    def test_emf_embedded_pdf_is_rasterized_before_vision(self):
        import struct
        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page(width=300, height=200)
        page.insert_text((20, 60), "Visible title", fontsize=24)
        pdf_bytes = pdf.tobytes()
        pdf.close()

        comment_data = b"GDIC" + pdf_bytes
        padding = bytes(-(12 + len(comment_data)) % 4)
        comment_size = 12 + len(comment_data) + len(padding)
        comment = (
            struct.pack("<III", 70, comment_size, len(comment_data))
            + comment_data + padding
        )
        end = struct.pack("<IIii", 14, 16, 0, 0)
        header = bytearray(108)
        total_size = len(header) + len(comment) + len(end)
        struct.pack_into("<II", header, 0, 1, len(header))
        struct.pack_into("<4i", header, 8, 0, 0, 300, 200)
        struct.pack_into("<4i", header, 24, 0, 0, 7620, 5080)
        struct.pack_into("<5I", header, 40, 0x464D4520, 0x10000, total_size, 3, 1)
        emf_data = bytes(header) + comment + end

        request = Mock()
        request.results.return_value = []
        handler = Mock()
        handler.performRequests_error_.return_value = (True, None)
        vision = SimpleNamespace(
            VNRecognizeTextRequest=Mock(),
            VNRequestTextRecognitionLevelAccurate=1,
            VNRequestTextRecognitionLevelFast=2,
            VNImageRequestHandler=Mock(),
        )
        vision.VNRecognizeTextRequest.alloc.return_value.init.return_value = request
        vision.VNImageRequestHandler.alloc.return_value.initWithData_options_.return_value = handler
        captured_image = {}

        def make_nsdata(data, length):
            captured_image["data"] = bytes(data)
            self.assertEqual(length, len(data))
            return object()

        foundation = SimpleNamespace(
            NSData=SimpleNamespace(dataWithBytes_length_=make_nsdata)
        )
        with (patch.object(sys, "platform", "darwin"),
              patch.dict(sys.modules, {
                  "Vision": vision, "Foundation": foundation
              })):
            self.assertEqual(
                self.namespace["_md_vision_ocr"](
                    emf_data, context="PowerPoint slide 24"
                ),
                "",
            )

        rasterized = captured_image["data"]
        self.assertTrue(rasterized.startswith(b"\x89PNG\r\n\x1a\n"))
        from PIL import Image

        with Image.open(io.BytesIO(rasterized)) as image:
            self.assertGreater(image.width, 300)
            self.assertGreater(image.height, 200)

    def test_failed_emf_ocr_image_is_saved_as_rasterized_png(self):
        import struct
        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page(width=100, height=80)
        page.insert_text((12, 30), "Visible title", fontsize=18)
        pdf_bytes = pdf.tobytes()
        pdf.close()

        comment_data = b"GDIC" + pdf_bytes
        padding = bytes(-(12 + len(comment_data)) % 4)
        comment_size = 12 + len(comment_data) + len(padding)
        comment = (
            struct.pack("<III", 70, comment_size, len(comment_data))
            + comment_data + padding
        )
        end = struct.pack("<IIii", 14, 16, 0, 0)
        header = bytearray(108)
        total_size = len(header) + len(comment) + len(end)
        struct.pack_into("<II", header, 0, 1, len(header))
        struct.pack_into("<4i", header, 8, 0, 0, 100, 80)
        struct.pack_into("<4i", header, 24, 0, 0, 2540, 2032)
        struct.pack_into("<5I", header, 40, 0x464D4520, 0x10000, total_size, 3, 1)
        emf_data = bytes(header) + comment + end

        handler = Mock()
        handler.performRequests_error_.return_value = (False, "OCR failed")
        vision = SimpleNamespace(
            VNRecognizeTextRequest=Mock(),
            VNRequestTextRecognitionLevelAccurate=1,
            VNRequestTextRecognitionLevelFast=2,
            VNImageRequestHandler=Mock(),
        )
        vision.VNRecognizeTextRequest.alloc.return_value.init.return_value = Mock()
        vision.VNImageRequestHandler.alloc.return_value.initWithData_options_.return_value = handler
        foundation = SimpleNamespace(
            NSData=SimpleNamespace(
                dataWithBytes_length_=Mock(return_value=object())
            )
        )
        failed_dir = self.root / "failed-images"
        with (patch.object(sys, "platform", "darwin"),
              patch.dict(sys.modules, {
                  "Vision": vision, "Foundation": foundation
              })):
            with self.assertRaisesRegex(OSError, "Apple Vision OCR failed"):
                self.namespace["_md_vision_ocr"](
                    emf_data,
                    context="PowerPoint slide 24",
                    failed_dir=failed_dir,
                    failed_name="deck-slide-24-Picture 6",
                    image_extension=".wmf",
                )

        saved_image = failed_dir / "deck-slide-24-Picture_6.png"
        self.assertTrue(saved_image.read_bytes().startswith(b"\x89PNG\r\n\x1a\n"))
        from PIL import Image

        with Image.open(saved_image) as image:
            self.assertGreater(image.width, 100)
            self.assertGreater(image.height, 80)

    def test_docx_ocr_is_inserted_after_embedded_image(self):
        from docx import Document
        from docx.shared import Inches

        source = self.root / "document.docx"
        document = Document()
        document.add_paragraph().add_run().add_picture(str(self.image_path), width=Inches(1))
        document.save(source)
        recognize = Mock(return_value="Word image text")
        failed_dir = self.root / "failed-images"
        with patch.dict(self.namespace, {
                "_md_require_vision": Mock(),
                "_md_vision_ocr": recognize,
        }), redirect_stdout(io.StringIO()) as output:
            markdown = self.namespace["docx2md"](
                source, ocr=True, ocr_failed_dir=failed_dir, overwrite=True
            )
        self.assertEqual(output.getvalue(), markdown)
        self.assertIn("Word image text", markdown)
        self.assertIn("\n[Image OCR]", markdown)
        recognize.assert_called_once()
        self.assertIn(str(source), recognize.call_args.args[3])
        self.assertEqual(recognize.call_args.kwargs["failed_dir"], str(failed_dir))
        self.assertTrue(recognize.call_args.kwargs["overwrite"])

    def test_xlsx_ocr_includes_image_anchor_and_sheet_context(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SpreadsheetImage

        source = self.root / "workbook.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Summary"
        worksheet["A1"] = "Report"
        worksheet.add_image(SpreadsheetImage(str(self.image_path)), "B2")
        workbook.save(source)
        recognize = Mock(return_value="Spreadsheet image text")
        failed_dir = self.root / "failed-images"
        with patch.dict(self.namespace, {
                "_md_require_vision": Mock(),
                "_md_vision_ocr": recognize,
        }), redirect_stdout(io.StringIO()) as output:
            markdown = self.namespace["xlsx2md"](
                source, ocr=True, ocr_failed_dir=failed_dir, overwrite=True
            )
        self.assertEqual(output.getvalue(), markdown)
        self.assertIn("## Summary", markdown)
        self.assertIn("### OCR from image near B2", markdown)
        self.assertIn("Spreadsheet image text", markdown)
        recognize.assert_called_once()
        self.assertIn(str(source), recognize.call_args.args[3])
        self.assertEqual(recognize.call_args.kwargs["failed_dir"], str(failed_dir))
        self.assertTrue(recognize.call_args.kwargs["overwrite"])

    def test_xlsx_ocr_skips_unreadable_image_with_warning(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SpreadsheetImage
        from zipfile import ZIP_DEFLATED, ZipFile
        from PIL import Image

        image_path = self.root / "image.bmp"
        Image.new("RGB", (32, 32), "white").save(image_path)
        source = self.root / "workbook.xlsx"
        workbook = Workbook()
        workbook.active.add_image(SpreadsheetImage(str(image_path)), "B2")
        workbook.save(source)

        broken_source = self.root / "broken.xlsx"
        with ZipFile(source) as original, ZipFile(
                broken_source, "w", ZIP_DEFLATED) as broken:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == "xl/media/image1.bmp":
                    data = image_path.read_bytes()[:60]
                broken.writestr(item, data)

        recognize = Mock()
        failed_dir = self.root / "failed-images"
        failed_dir.mkdir()
        existing_failed_image = failed_dir / "broken-Sheet-B2.bmp"
        existing_failed_image.write_bytes(b"old image")
        with (patch.dict(self.namespace, {
                "_md_check_ocr_platform": Mock(),
                "_md_vision_ocr": recognize,
        }), warnings.catch_warnings(record=True) as caught,
                redirect_stdout(io.StringIO())):
            warnings.simplefilter("always")
            markdown = self.namespace["xlsx2md"](
                broken_source, ocr=True, ocr_failed_dir=failed_dir,
                overwrite=True,
            )
        self.assertIn("## Sheet", markdown)
        self.assertEqual(len(caught), 1)
        self.assertIn(str(broken_source), str(caught[0].message))
        self.assertIn("worksheet 'Sheet', image near B2", str(caught[0].message))
        self.assertIn("invalid or unreadable", str(caught[0].message))
        saved_image = failed_dir / "broken-Sheet-B2.bmp"
        self.assertEqual(saved_image.read_bytes(), image_path.read_bytes()[:60])
        self.assertIn(str(saved_image), str(caught[0].message))
        recognize.assert_not_called()


if __name__ == "__main__":
    unittest.main()
